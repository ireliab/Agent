"""Generates the input files the fixtures attach.

They are built rather than committed so the expected answers live next to the
file that contains them - if you change a number here, the fixture that checks
for it is two files away, not in a binary nobody can read.

The PDF and PNG are written byte by byte because the project has no PDF writer
or imaging library, and neither is worth adding just to make test inputs.
"""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

ASSET_DIR = Path(__file__).parent / "assets"

# Facts deliberately chosen to be unguessable, so an answer containing one is
# evidence the model actually read the file rather than inferring from context.
PDF_NAME = "quarterly-notes.pdf"
PDF_FACT = "Northwind"
PDF_NUMBER = "41,920"

XLSX_NAME = "sales-figures.xlsx"
XLSX_TOTAL = 11700  # 2400 + 3100 + 1850 + 4350

CJK_NAME = "營運摘要_第2頁.docx"
CJK_FACT = "Kestrel"

IMAGE_NAME = "colour-bands.png"
IMAGE_COLOURS = ("red", "green", "blue")


def _pdf_escape(text: str) -> str:
    return text.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")


def build_pdf(path: Path) -> None:
    """A one-page PDF with selectable text, assembled by hand.

    Offsets in the cross-reference table have to be byte-exact, so the objects
    are serialised first and their positions recorded as they are appended.
    """
    lines = [
        "Quarterly Operating Notes",
        "",
        f"Prepared for the {PDF_FACT} division.",
        f"Total units shipped in the quarter: {PDF_NUMBER}.",
        "The warehouse relocation completed ahead of schedule.",
    ]
    text_ops = ["BT", "/F1 12 Tf", "72 720 Td", "14 TL"]
    for line in lines:
        text_ops.append(f"({_pdf_escape(line)}) Tj T*")
    text_ops.append("ET")
    stream = "\n".join(text_ops).encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_at = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n".encode()
    out += f"startxref\n{xref_at}\n%%EOF\n".encode()

    path.write_bytes(bytes(out))


def _png_chunk(kind: bytes, data: bytes) -> bytes:
    return (
        struct.pack(">I", len(data))
        + kind
        + data
        + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    )


def build_png(path: Path, width: int = 240, height: int = 240) -> None:
    """Three horizontal bands: red, green, blue, top to bottom.

    Colour is the most robustly checkable thing a vision tower can report - no
    OCR, no object naming, no ambiguity about what the right answer is.
    """
    bands = [(220, 30, 30), (30, 170, 60), (40, 80, 210)]
    raw = bytearray()
    for y in range(height):
        raw.append(0)  # filter type 0 (None) for this scanline
        red, green, blue = bands[min(y * len(bands) // height, len(bands) - 1)]
        raw += bytes((red, green, blue)) * width

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)  # 8-bit truecolour
    path.write_bytes(
        b"\x89PNG\r\n\x1a\n"
        + _png_chunk(b"IHDR", header)
        + _png_chunk(b"IDAT", zlib.compress(bytes(raw), 9))
        + _png_chunk(b"IEND", b"")
    )


def build_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "Q3"
    sheet.append(["Region", "Units"])
    for region, units in [
        ("North", 2400),
        ("South", 3100),
        ("East", 1850),
        ("West", 4350),
    ]:
        sheet.append([region, units])
    book.save(path)


def build_cjk_docx(path: Path) -> None:
    """A document whose filename a small model cannot retype correctly.

    This is the regression test for a real failure: the model re-tokenised
    `第2頁` as `第2 頁`, exact-match lookup failed, and it retried the same
    broken name until the run budget ran out.
    """
    from docx import Document

    document = Document()
    document.add_heading("營運摘要", level=1)
    document.add_paragraph(f"本季度的主要客戶為 {CJK_FACT} Holdings。")
    document.add_paragraph("The primary client this quarter was " f"{CJK_FACT} Holdings.")
    document.save(path)


BUILDERS = {
    PDF_NAME: build_pdf,
    IMAGE_NAME: build_png,
    XLSX_NAME: build_xlsx,
    CJK_NAME: build_cjk_docx,
}


def ensure_assets(force: bool = False) -> Path:
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    for name, builder in BUILDERS.items():
        path = ASSET_DIR / name
        if force or not path.exists():
            builder(path)
    return ASSET_DIR


if __name__ == "__main__":
    directory = ensure_assets(force=True)
    for asset in sorted(directory.iterdir()):
        print(f"{asset.name:28} {asset.stat().st_size:>8} bytes")
