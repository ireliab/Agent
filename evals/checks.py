"""Checks a fixture can assert about a run.

Each check is a small object with a name and a `run(result)` returning
(passed, detail). Keep them mechanical: a check that needs an opinion belongs
in `judged()`, which records the question instead of pretending to answer it.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

# A check returns (passed, detail). `passed = None` means "not scored" - used by
# judged(), so a subjective item shows up in the report without silently
# counting as a pass.
Outcome = tuple[bool | None, str]


@dataclass
class Check:
    name: str
    run: Callable[[Any], Outcome]


def _norm(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").casefold()


def calls_tool(name: str, at_least: int = 1, at_most: int | None = None) -> Check:
    """The agent must use this tool (counting only calls that actually ran)."""

    def run(result) -> Outcome:
        ran = [c for c in result.tool_calls if c["name"] == name and c["status"] == "ok"]
        count = len(ran)
        if count < at_least:
            names = sorted({c["name"] for c in result.tool_calls}) or ["(none)"]
            return False, f"ran {count}x, wanted >={at_least}; called: {', '.join(names)}"
        if at_most is not None and count > at_most:
            return False, f"ran {count}x, wanted <={at_most}"
        return True, f"{count}x"

    bound = f">={at_least}" + (f",<={at_most}" if at_most is not None else "")
    return Check(f"calls_tool:{name}{bound}", run)


def requests_tool(name: str, decision: str | None = None) -> Check:
    """The agent asked to use this tool, whether or not it was allowed to.

    `calls_tool` only sees calls that executed: a tool stopped at the approval
    step never reaches the tool wrapper. Use this one when what matters is that
    the agent reached for the tool, not that the tool ran.
    """

    def run(result) -> Outcome:
        asked = [a for a in result.approvals if a["name"] == name]
        if not asked:
            seen = sorted({a["name"] for a in result.approvals}) or ["(none)"]
            return False, f"not requested; requested: {', '.join(seen)}"
        if decision and not any(a["decision"] == decision for a in asked):
            return False, f"requested but decided {asked[0]['decision']}, wanted {decision}"
        return True, ", ".join(f"{a['name']}:{a['decision']}" for a in asked)

    suffix = f" ({decision})" if decision else ""
    return Check(f"requests_tool:{name}{suffix}", run)


def never_calls_tool(name: str) -> Check:
    def run(result) -> Outcome:
        count = sum(1 for c in result.tool_calls if c["name"] == name)
        return (count == 0), (f"called {count}x" if count else "not called")

    return Check(f"never_calls_tool:{name}", run)


def answer_contains(*needles: str, any_of: bool = False) -> Check:
    """Substrings the final answer must contain, ignoring case and whitespace."""

    def run(result) -> Outcome:
        answer = _norm(result.answer)
        hits = [n for n in needles if _norm(n) in answer]
        ok = bool(hits) if any_of else len(hits) == len(needles)
        missing = [n for n in needles if n not in hits]
        return ok, ("" if ok else f"missing: {', '.join(missing)}")

    joiner = " | " if any_of else " & "
    return Check(f"answer_contains:{joiner.join(needles)}", run)


def answer_excludes(*needles: str) -> Check:
    def run(result) -> Outcome:
        answer = _norm(result.answer)
        found = [n for n in needles if _norm(n) in answer]
        return (not found), (f"found: {', '.join(found)}" if found else "")

    return Check(f"answer_excludes:{', '.join(needles)}", run)


def answer_matches(pattern: str, flags: int = re.I) -> Check:
    def run(result) -> Outcome:
        match = re.search(pattern, result.answer or "", flags)
        return bool(match), (match.group(0)[:80] if match else "no match")

    return Check(f"answer_matches:{pattern}", run)


def answer_not_empty(min_chars: int = 20) -> Check:
    def run(result) -> Outcome:
        count = len((result.answer or "").strip())
        return count >= min_chars, f"{count} chars"

    return Check(f"answer_not_empty:>={min_chars}", run)


def answer_shorter_than(max_chars: int) -> Check:
    """Brevity when brevity was asked for. Small models pad; this catches it."""

    def run(result) -> Outcome:
        count = len((result.answer or "").strip())
        return count <= max_chars, f"{count} chars (max {max_chars})"

    return Check(f"answer_shorter_than:{max_chars}", run)


def answer_order(*needles: str) -> Check:
    """The needles must all appear, in this order."""

    def run(result) -> Outcome:
        answer = _norm(result.answer)
        position = -1
        for needle in needles:
            found = answer.find(_norm(needle), position + 1)
            if found < 0:
                return False, f"{needle} not found after position {position}"
            position = found
        return True, "in order"

    return Check(f"answer_order:{' -> '.join(needles)}", run)


def file_written(suffix: str = "", at_least: int = 1) -> Check:
    """A real file must exist in this run's output directory."""

    def run(result) -> Outcome:
        files = [f for f in result.output_files if f.lower().endswith(suffix.lower())]
        return len(files) >= at_least, ", ".join(files) or "(none)"

    return Check(f"file_written:{suffix or 'any'}", run)


def no_file_written() -> Check:
    def run(result) -> Outcome:
        return (not result.output_files), (", ".join(result.output_files) or "none")

    return Check("no_file_written", run)


def file_named(name: str) -> Check:
    def run(result) -> Outcome:
        ok = any(f.casefold() == name.casefold() for f in result.output_files)
        return ok, ", ".join(result.output_files) or "(none)"

    return Check(f"file_named:{name}", run)


def no_file_named(name: str) -> Check:
    """Nothing by this name was written - e.g. an edited filename replaced it."""

    def run(result) -> Outcome:
        hit = [f for f in result.output_files if f.casefold() == name.casefold()]
        return (not hit), (f"found {name}" if hit else "absent")

    return Check(f"no_file_named:{name}", run)


def docx_contains(
    *needles: str,
    headings: int = 0,
    tables: int = 0,
    table_rows: int = 0,
) -> Check:
    """Open the produced .docx and assert its real structure, not the model's claim."""

    def run(result) -> Outcome:
        from docx import Document

        names = [f for f in result.output_files if f.endswith(".docx")]
        if not names:
            return False, "no .docx produced"
        document = Document(str(Path(result.output_dir) / names[0]))
        text = _norm("\n".join(p.text for p in document.paragraphs))
        for table in document.tables:
            text += " " + _norm(" ".join(c.text for row in table.rows for c in row.cells))

        problems = []
        missing = [n for n in needles if _norm(n) not in text]
        if missing:
            problems.append(f"missing text: {', '.join(missing)}")
        found = sum(1 for p in document.paragraphs if p.style.name.startswith("Heading"))
        if found < headings:
            problems.append(f"{found} headings, wanted >={headings}")
        if len(document.tables) < tables:
            problems.append(f"{len(document.tables)} tables, wanted >={tables}")
        if tables and document.tables and len(document.tables[0].rows) < table_rows:
            problems.append(f"{len(document.tables[0].rows)} rows, wanted >={table_rows}")
        summary = f"{found} headings, {len(document.tables)} tables"
        return (not problems), "; ".join(problems) or summary

    return Check(f"docx_contains:{', '.join(needles) or 'structure'}", run)


def max_model_calls(limit: int) -> Check:
    """Efficiency. A model that needs twelve calls for a one-step task is a finding."""

    def run(result) -> Outcome:
        return result.model_calls <= limit, f"{result.model_calls} calls (limit {limit})"

    return Check(f"max_model_calls:{limit}", run)


def terminates_cleanly() -> Check:
    def run(result) -> Outcome:
        return result.termination == "normal", result.termination

    return Check("terminates_cleanly", run)


def no_blocked_calls() -> Check:
    """The repeated-call guard never fired, i.e. the agent did not loop."""

    def run(result) -> Outcome:
        blocked = [c["name"] for c in result.tool_calls if c["status"] == "blocked"]
        return (not blocked), (f"blocked: {', '.join(blocked)}" if blocked else "none")

    return Check("no_blocked_calls", run)


def no_failed_calls() -> Check:
    def run(result) -> Outcome:
        bad = [
            f"{c['name']}({c['status']})"
            for c in result.tool_calls
            if c["status"] in {"error", "exception"}
        ]
        return (not bad), (", ".join(bad) if bad else "none")

    return Check("no_failed_calls", run)


def tool_arg_matches(tool: str, arg: str, pattern: str) -> Check:
    """Assert what the model actually put in a tool argument."""

    def run(result) -> Outcome:
        values = [str(c["args"].get(arg, "")) for c in result.tool_calls if c["name"] == tool]
        if not values:
            return False, f"{tool} never called"
        hits = [v for v in values if re.search(pattern, v, re.I)]
        return bool(hits), "saw: " + ", ".join(v[:60] for v in values)

    return Check(f"tool_arg_matches:{tool}.{arg}", run)


def judged(question: str) -> Check:
    """Record a question that needs a human eye. Never counts as a pass or a fail.

    Anything checkable mechanically should be checked mechanically. This exists
    so the subjective remainder is visible in the report rather than omitted.
    """

    def run(_result) -> Outcome:
        return None, question

    return Check(f"judge:{question[:60]}", run)
