"""Runs one fixture against the real agent and measures what happened.

The agent is driven exactly as the web server drives it - same `build_agent`,
same `build_user_content`, same streaming modes - so a result here means the
same thing as a result in the browser. The only differences are deliberate:
each run gets a fresh in-memory checkpointer and its own output and upload
directories, so runs cannot see each other's files or collide on a filename.
"""

from __future__ import annotations

import asyncio
import json
import shutil
import time
import uuid
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any

from langchain_core.messages import AIMessage, AIMessageChunk
from langgraph.checkpoint.memory import InMemorySaver
from langgraph.errors import GraphRecursionError
from langgraph.types import Command

from evals.probe import Probe
from testing.agent import RECURSION_LIMIT, build_agent, build_user_content
from testing.tools import set_file_roots, set_upload_scope, upload_dir

# Text langchain injects when a limit stops the run. These messages are added to
# state rather than generated, so the only way to notice one is to read it.
MODEL_LIMIT_MARKER = "Model call limits exceeded"
TOOL_LIMIT_MARKER = "call limit reached"

DEFAULT_TIMEOUT_S = 300.0


@dataclass
class RunResult:
    fixture: str
    rep: int
    started_at: str
    ok: bool = True
    error: str | None = None
    answer: str = ""
    termination: str = "normal"
    model_calls: int = 0
    # Only calls that actually executed. A call refused at the approval step
    # never reaches the tool wrapper, so it appears in `approvals` instead.
    tool_calls: list[dict] = field(default_factory=list)
    approvals: list[dict] = field(default_factory=list)
    latency: dict[str, float] = field(default_factory=dict)
    tokens: dict[str, int] = field(default_factory=dict)
    output_dir: str = ""
    output_files: list[str] = field(default_factory=list)
    artifact_excerpt: str = ""
    checks: list[dict] = field(default_factory=list)
    passed: bool = False

    def key(self) -> str:
        return f"{self.fixture}#{self.rep}"


def _text_of(message: Any) -> str:
    """Flatten message content to text, ignoring image and tool-call blocks."""
    content = getattr(message, "content", "")
    if isinstance(content, str):
        return content
    parts = []
    for block in content or []:
        if isinstance(block, str):
            parts.append(block)
        elif isinstance(block, dict) and block.get("type") == "text":
            parts.append(block.get("text", ""))
    return "".join(parts)


def _decisions_for(requests: list[dict], policy: dict[str, Any]) -> list[dict]:
    """Turn a fixture's approval policy into one decision per pending call.

    langchain requires exactly one decision per request, in order, so this must
    answer every request even when the policy only cares about one of them.
    """
    decision_type = policy.get("decision", "approve")
    decisions = []
    for request in requests:
        if decision_type == "edit":
            args = dict(request.get("args") or {})
            args.update(policy.get("edit_args") or {})
            decisions.append(
                {
                    "type": "edit",
                    "edited_action": {"name": request.get("name"), "args": args},
                }
            )
        elif decision_type in {"reject", "respond"}:
            decisions.append(
                {"type": decision_type, "message": policy.get("message", "Not this time.")}
            )
        else:
            decisions.append({"type": "approve"})
    return decisions


class _Trace:
    """Accumulates timings and events across the streaming passes of one run."""

    def __init__(self) -> None:
        self.started = time.perf_counter()
        self.first_chunk_s: float | None = None
        self.first_text_s: float | None = None
        self.input_tokens = 0
        self.output_tokens = 0
        self.saw_interrupt = False
        self.interrupt_requests: list[dict] = []

    def note_chunk(self, message: AIMessageChunk) -> None:
        now = time.perf_counter() - self.started
        if self.first_chunk_s is None:
            self.first_chunk_s = now
        if self.first_text_s is None and _text_of(message).strip():
            self.first_text_s = now
        if message.usage_metadata:
            self.input_tokens += message.usage_metadata.get("input_tokens", 0)
            self.output_tokens += message.usage_metadata.get("output_tokens", 0)


async def _drive(agent, agent_input: Any, config: dict, trace: _Trace) -> None:
    """Stream one pass of the agent, stopping if it pauses for approval."""
    trace.saw_interrupt = False
    trace.interrupt_requests = []
    async for mode, chunk in agent.astream(
        agent_input, config=config, stream_mode=["messages", "updates"]
    ):
        if mode == "messages":
            message, _meta = chunk
            if isinstance(message, AIMessageChunk):
                trace.note_chunk(message)
        elif mode == "updates":
            for node, update in (chunk or {}).items():
                if node == "__interrupt__":
                    trace.saw_interrupt = True
                    for interrupt in update or ():
                        value = getattr(interrupt, "value", {}) or {}
                        trace.interrupt_requests.extend(value.get("action_requests", []))


async def _execute(fixture, probe: Probe, trace: _Trace, thread_id: str):
    """Run every turn of the fixture, answering any approval it pauses on."""
    agent = build_agent(checkpointer=InMemorySaver(), extra_middleware=[probe])
    config = {
        "configurable": {"thread_id": thread_id},
        "recursion_limit": fixture.recursion_limit or RECURSION_LIMIT,
    }
    approvals: list[dict] = []

    for index, turn in enumerate(fixture.turns):
        # Attachments belong to the first turn, as they would in the browser.
        content = build_user_content(
            turn, fixture.attachments if index == 0 else [], thread_id
        )
        agent_input: Any = {"messages": [{"role": "user", "content": content}]}

        # A turn can pause more than once - one pause per batch of guarded tools.
        for _ in range(fixture.max_approval_rounds):
            await _drive(agent, agent_input, config, trace)
            if not trace.saw_interrupt:
                break
            policy = fixture.approval or {"decision": "approve"}
            requests = list(trace.interrupt_requests)
            decisions = _decisions_for(requests, policy)
            approvals.extend(
                {
                    "name": request.get("name"),
                    "args": request.get("args") or {},
                    "decision": decision["type"],
                }
                for request, decision in zip(requests, decisions)
            )
            agent_input = Command(resume={"decisions": decisions})
        else:
            raise RuntimeError(
                f"still waiting for approval after {fixture.max_approval_rounds} rounds"
            )

    snapshot = await agent.aget_state(config)
    return snapshot.values.get("messages") or [], approvals


def _excerpt(outputs: Path, names: list[str], limit: int = 700) -> str:
    """Read back what the run produced, so a human can judge the artifact itself.

    Judging "is the document any good?" against the model's closing remark is
    judging the wrong thing - it says the file is fine no matter what is in it.
    """
    for name in names:
        path = outputs / name
        try:
            if name.endswith(".docx"):
                from docx import Document

                document = Document(str(path))
                text = "\n".join(p.text for p in document.paragraphs if p.text.strip())
                for table in document.tables:
                    for row in table.rows:
                        text += "\n| " + " | ".join(c.text for c in row.cells) + " |"
            elif name.endswith((".txt", ".md", ".csv")):
                text = path.read_text(encoding="utf-8", errors="replace")
            else:
                continue
        except Exception as exc:
            return f"(could not read {name}: {type(exc).__name__})"
        text = text.strip()
        return text[:limit] + (" ..." if len(text) > limit else "")
    return ""


def _summarise(messages: list[Any]) -> tuple[str, str]:
    """Pull the final answer out of state, and work out how the run ended."""
    termination = "normal"
    answer = ""
    for message in messages:
        if not isinstance(message, AIMessage):
            continue
        text = _text_of(message)
        if MODEL_LIMIT_MARKER in text:
            termination = "model_limit"
            continue
        if TOOL_LIMIT_MARKER in text:
            termination = "tool_limit"
            continue
        if text.strip():
            answer = text
    return answer.strip(), termination


async def run_fixture(fixture, rep: int, root: Path) -> RunResult:
    """Run one fixture once, in its own directories, and score it.

    Context variables are set inside this coroutine on purpose: asyncio gives
    each task its own copy of the context, so concurrent runs stay isolated.
    """
    started_at = time.strftime("%Y-%m-%dT%H:%M:%S")
    scratch = root / fixture.name / str(rep)
    outputs = scratch / "outputs"
    uploads = scratch / "uploads"
    if scratch.exists():
        shutil.rmtree(scratch)
    outputs.mkdir(parents=True)
    uploads.mkdir(parents=True)

    thread_id = f"eval-{fixture.name}-{rep}-{uuid.uuid4().hex[:8]}"
    set_file_roots(outputs=outputs, uploads=uploads)
    set_upload_scope(thread_id)

    if fixture.attachments:
        target = upload_dir(thread_id)
        target.mkdir(parents=True, exist_ok=True)
        for name in fixture.attachments:
            source = fixture.asset_dir / name
            if not source.exists():
                raise FileNotFoundError(f"fixture asset missing: {source}")
            shutil.copy2(source, target / name)

    result = RunResult(
        fixture=fixture.name, rep=rep, started_at=started_at, output_dir=str(outputs)
    )
    probe = Probe()
    trace = _Trace()
    wall = time.perf_counter()

    try:
        messages, approvals = await asyncio.wait_for(
            _execute(fixture, probe, trace, thread_id),
            timeout=fixture.timeout_s,
        )
        result.answer, result.termination = _summarise(messages)
        result.approvals = approvals
    except asyncio.TimeoutError:
        result.ok, result.termination = False, "timeout"
        result.error = f"exceeded {fixture.timeout_s}s"
    except GraphRecursionError as exc:
        result.ok, result.termination = False, "recursion"
        result.error = str(exc)[:300]
    except Exception as exc:  # a harness or transport failure, not a model verdict
        result.ok, result.termination = False, "error"
        result.error = f"{type(exc).__name__}: {exc}"[:300]

    total = time.perf_counter() - wall
    tool_s = sum(c.duration_s for c in probe.tool_calls)
    model_s = sum(c.duration_s for c in probe.model_calls)
    result.model_calls = len(probe.model_calls)
    result.tool_calls = [asdict(c) for c in probe.tool_calls]
    result.latency = {
        "total_s": round(total, 3),
        "ttft_s": round(trace.first_chunk_s, 3) if trace.first_chunk_s else None,
        "first_text_s": round(trace.first_text_s, 3) if trace.first_text_s else None,
        "model_s": round(model_s, 3),
        "tool_s": round(tool_s, 3),
        # Whatever is left is the graph itself: middleware, state writes, checkpointing.
        "overhead_s": round(total - model_s - tool_s, 3),
    }
    result.tokens = {
        "input": trace.input_tokens,
        "output": trace.output_tokens,
        # Honest denominator: time actually spent inside model calls.
        "output_tps": round(trace.output_tokens / model_s, 1) if model_s > 0 else None,
    }
    result.output_files = sorted(p.name for p in outputs.iterdir() if p.is_file())
    result.artifact_excerpt = _excerpt(outputs, result.output_files)

    scored = []
    for check in fixture.checks:
        try:
            passed, detail = check.run(result)
        except Exception as exc:
            passed, detail = False, f"check raised {type(exc).__name__}: {exc}"
        scored.append({"name": check.name, "passed": passed, "detail": detail})
    result.checks = scored
    # Unscored checks (the judged ones) do not count either way.
    result.passed = result.ok and all(c["passed"] is not False for c in scored)
    return result


def append_jsonl(path: Path, result: RunResult) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(asdict(result), ensure_ascii=False, default=str) + "\n")


def completed_keys(path: Path) -> set[str]:
    """Which (fixture, rep) pairs already have a result, so a run can resume."""
    if not path.exists():
        return set()
    done = set()
    for line in path.read_text(encoding="utf-8").splitlines():
        if not line.strip():
            continue
        try:
            row = json.loads(line)
        except json.JSONDecodeError:
            continue  # a half-written line from an interrupted run
        done.add(f"{row['fixture']}#{row['rep']}")
    return done
