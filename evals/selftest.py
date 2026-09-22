"""Tests the ruler, not the thing being measured.

A check that always returns True is invisible: the suite goes green and tells
you nothing. So every check is exercised twice here, once against a result it
should accept and once against a result it should reject, and it has to get
both right.

    python -m evals.selftest
"""

from __future__ import annotations

import shutil
import tempfile
from pathlib import Path

from docx import Document

from evals import checks
from evals.runner import RunResult, _decisions_for, _summarise
from langchain_core.messages import AIMessage, HumanMessage


def make_result(**overrides) -> RunResult:
    base = {
        "fixture": "synthetic",
        "rep": 0,
        "started_at": "now",
        "answer": "The Northwind division shipped 41,920 units.",
        "termination": "normal",
        "model_calls": 3,
        "tool_calls": [
            {"name": "read_document", "args": {"filename": "1"}, "status": "ok", "duration_s": 0.1},
        ],
        "output_files": [],
        "output_dir": "",
    }
    base.update(overrides)
    return RunResult(**base)


def expect(check: checks.Check, result: RunResult, want: bool | None, label: str) -> None:
    got, detail = check.run(result)
    if got is not want:
        raise AssertionError(
            f"{label}: {check.name} returned {got!r}, expected {want!r} (detail: {detail})"
        )


def check_pairs(tmp: Path) -> None:
    good = make_result()

    # -- tool checks ---------------------------------------------------
    expect(checks.calls_tool("read_document"), good, True, "calls_tool accepts")
    expect(checks.calls_tool("write_docx"), good, False, "calls_tool rejects missing")
    expect(
        checks.calls_tool("read_document", at_most=0), good, False, "calls_tool rejects too many"
    )
    # A call that ran but errored must not count as having used the tool.
    errored = make_result(
        tool_calls=[{"name": "read_document", "args": {}, "status": "error", "duration_s": 0.1}]
    )
    expect(checks.calls_tool("read_document"), errored, False, "calls_tool ignores failed calls")

    expect(checks.never_calls_tool("write_docx"), good, True, "never_calls_tool accepts")
    expect(checks.never_calls_tool("read_document"), good, False, "never_calls_tool rejects")

    blocked = make_result(
        tool_calls=[
            {
                "name": "read_document",
                "args": {},
                "status": "blocked",
                "duration_s": 0.0,
                "detail": "Blocked: ...",
            }
        ]
    )
    expect(checks.no_blocked_calls(), good, True, "no_blocked_calls accepts")
    expect(checks.no_blocked_calls(), blocked, False, "no_blocked_calls rejects")
    expect(checks.no_failed_calls(), good, True, "no_failed_calls accepts")
    expect(checks.no_failed_calls(), errored, False, "no_failed_calls rejects")

    expect(
        checks.tool_arg_matches("read_document", "filename", r"^\d+$"),
        good,
        True,
        "tool_arg_matches accepts",
    )
    expect(
        checks.tool_arg_matches("read_document", "filename", r"budget"),
        good,
        False,
        "tool_arg_matches rejects",
    )

    # -- answer checks -------------------------------------------------
    expect(checks.answer_contains("Northwind"), good, True, "answer_contains accepts")
    expect(checks.answer_contains("Northwind", "Kestrel"), good, False, "answer_contains rejects")
    expect(
        checks.answer_contains("Northwind", "Kestrel", any_of=True),
        good,
        True,
        "answer_contains any_of accepts",
    )
    expect(checks.answer_excludes("Kestrel"), good, True, "answer_excludes accepts")
    expect(checks.answer_excludes("Northwind"), good, False, "answer_excludes rejects")
    expect(checks.answer_matches(r"41[,\s]?920"), good, True, "answer_matches accepts")
    expect(checks.answer_matches(r"99[,\s]?999"), good, False, "answer_matches rejects")
    expect(checks.answer_not_empty(10), good, True, "answer_not_empty accepts")
    expect(checks.answer_not_empty(10), make_result(answer="hi"), False, "answer_not_empty rejects")
    expect(checks.answer_shorter_than(100), good, True, "answer_shorter_than accepts")
    expect(checks.answer_shorter_than(5), good, False, "answer_shorter_than rejects")

    ordered = make_result(answer="red at the top, green in the middle, blue at the bottom")
    expect(checks.answer_order("red", "green", "blue"), ordered, True, "answer_order accepts")
    expect(checks.answer_order("blue", "green", "red"), ordered, False, "answer_order rejects")
    # Not just presence - the same words in the wrong order must fail.
    expect(
        checks.answer_order("red", "green", "blue"),
        make_result(answer="blue, green, red"),
        False,
        "answer_order is not just containment",
    )

    # -- run-shape checks ----------------------------------------------
    expect(checks.max_model_calls(5), good, True, "max_model_calls accepts")
    expect(checks.max_model_calls(2), good, False, "max_model_calls rejects")
    expect(checks.terminates_cleanly(), good, True, "terminates_cleanly accepts")
    expect(
        checks.terminates_cleanly(),
        make_result(termination="model_limit"),
        False,
        "terminates_cleanly rejects",
    )

    # -- file checks ---------------------------------------------------
    outputs = tmp / "outputs"
    outputs.mkdir(parents=True, exist_ok=True)
    document = Document()
    document.add_heading("Title", level=1)
    document.add_heading("Section", level=2)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Region"
    table.cell(1, 0).text = "North"
    document.save(outputs / "report.docx")

    with_file = make_result(output_files=["report.docx"], output_dir=str(outputs))
    expect(checks.file_written(".docx"), with_file, True, "file_written accepts")
    expect(checks.file_written(".pdf"), with_file, False, "file_written rejects")
    expect(checks.no_file_written(), good, True, "no_file_written accepts")
    expect(checks.no_file_written(), with_file, False, "no_file_written rejects")
    expect(checks.file_named("report.docx"), with_file, True, "file_named accepts")
    expect(checks.file_named("other.docx"), with_file, False, "file_named rejects")

    expect(checks.docx_contains("North"), with_file, True, "docx_contains reads tables")
    expect(checks.docx_contains("Nowhere"), with_file, False, "docx_contains rejects missing text")
    expect(checks.docx_contains(headings=2), with_file, True, "docx_contains counts headings")
    expect(
        checks.docx_contains(headings=5), with_file, False, "docx_contains rejects too few headings"
    )
    expect(checks.docx_contains(tables=1, table_rows=2), with_file, True, "docx_contains tables")
    expect(
        checks.docx_contains(tables=3), with_file, False, "docx_contains rejects too few tables"
    )
    expect(checks.docx_contains("x"), good, False, "docx_contains rejects when no docx exists")

    # -- judged is never a verdict -------------------------------------
    expect(checks.judged("is it good?"), good, None, "judged stays unscored")


def summary_pairs() -> None:
    """_summarise has to tell a real answer from an injected limit notice."""
    normal = [HumanMessage("hi"), AIMessage("Here is the answer.")]
    answer, how = _summarise(normal)
    assert (answer, how) == ("Here is the answer.", "normal"), (answer, how)

    limited = [
        HumanMessage("hi"),
        AIMessage("Working on it."),
        AIMessage("Model call limits exceeded: run limit (25/25)"),
    ]
    answer, how = _summarise(limited)
    assert how == "model_limit", how
    # The notice must not be mistaken for the agent's answer.
    assert answer == "Working on it.", answer

    tool_limited = [HumanMessage("hi"), AIMessage("'search' tool call limit reached: run limit.")]
    _, how = _summarise(tool_limited)
    assert how == "tool_limit", how


def decision_pairs() -> None:
    """Every pending call must get exactly one decision, in order."""
    requests = [
        {"name": "write_docx", "args": {"filename": "a.docx", "content": "x"}},
        {"name": "write_text_file", "args": {"filename": "b.md", "content": "y"}},
    ]

    approved = _decisions_for(requests, {"decision": "approve"})
    assert len(approved) == 2 and all(d["type"] == "approve" for d in approved), approved

    edited = _decisions_for(requests, {"decision": "edit", "edit_args": {"filename": "new.docx"}})
    assert len(edited) == 2, edited
    assert edited[0]["edited_action"]["args"]["filename"] == "new.docx", edited[0]
    # Editing one argument must not discard the others.
    assert edited[0]["edited_action"]["args"]["content"] == "x", edited[0]
    assert edited[0]["edited_action"]["name"] == "write_docx", edited[0]

    rejected = _decisions_for(requests, {"decision": "reject", "message": "no"})
    assert all(d["type"] == "reject" and d["message"] == "no" for d in rejected), rejected


def main() -> None:
    tmp = Path(tempfile.mkdtemp(prefix="eval-selftest-"))
    try:
        check_pairs(tmp)
        summary_pairs()
        decision_pairs()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    print("selftest OK - every check accepted a good result and rejected a bad one")


if __name__ == "__main__":
    main()
