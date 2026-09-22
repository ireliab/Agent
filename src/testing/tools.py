"""Tools that produce real files on disk.

The agent's built-in `write_file` tool writes into LangGraph state, which is a
virtual filesystem that never touches the disk. Anything the user should be able
to open afterwards has to be written here instead, into `OUTPUT_DIR`.
"""

import contextvars
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.shared import Pt

OUTPUT_DIR = Path.cwd() / "outputs"


def _safe_output_path(filename: str, suffix: str) -> Path:
    """Resolve a model-supplied filename to a path inside OUTPUT_DIR.

    The model chooses this name, so strip any directory part before using it.
    """
    stem = Path(filename).name.strip() or "report"
    if not stem.lower().endswith(suffix):
        stem = f"{Path(stem).stem}{suffix}"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    return OUTPUT_DIR / stem


def _add_runs(paragraph, text: str) -> None:
    """Add text to a paragraph, honouring **bold** and *italic* markers."""
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*\n]+\*)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("*") and part.endswith("*"):
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def write_docx(filename: str, content: str) -> str:
    """Save a report as a Word (.docx) file that the user can open.

    Args:
        filename: File name for the report, e.g. "quarterly-report.docx".
        content: The report body. Markdown is supported: `#` to `######` for
            headings, `-` or `*` for bullets, `1.` for numbered lists,
            `|` tables, and `**bold**` / `*italic*` inline.

    Returns:
        A confirmation message with the saved path.
    """
    path = _safe_output_path(filename, ".docx")
    document = Document()
    lines = content.split("\n")

    index = 0
    while index < len(lines):
        line = lines[index].strip()
        index += 1

        if not line:
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            document.add_heading(heading.group(2), level=min(len(heading.group(1)), 4))
            continue

        # table: a header row followed by a |---|---| separator
        next_line = lines[index].strip() if index < len(lines) else ""
        if "|" in line and "-" in next_line and re.fullmatch(r"[\s|:-]+", next_line):
            cells = lambda row: [c.strip() for c in row.strip().strip("|").split("|")]
            header = cells(line)
            index += 1
            rows = []
            while index < len(lines) and "|" in lines[index]:
                rows.append(cells(lines[index]))
                index += 1
            table = document.add_table(rows=1, cols=len(header))
            table.style = "Light Grid Accent 1"
            for cell, text in zip(table.rows[0].cells, header):
                cell.text = text
            for row in rows:
                target = table.add_row().cells
                for cell, text in zip(target, row):
                    cell.text = text
            document.add_paragraph()
            continue

        bullet = re.match(r"^[-*+]\s+(.*)$", line)
        if bullet:
            _add_runs(document.add_paragraph(style="List Bullet"), bullet.group(1))
            continue

        numbered = re.match(r"^\d+[.)]\s+(.*)$", line)
        if numbered:
            _add_runs(document.add_paragraph(style="List Number"), numbered.group(1))
            continue

        if re.fullmatch(r"(-{3,}|\*{3,}|_{3,})", line):
            continue

        _add_runs(document.add_paragraph(), line)

    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            if not run.font.size:
                run.font.size = Pt(11)

    document.save(path)
    return f"Saved to {path}. The user can download it from the Files panel."


def write_text_file(filename: str, content: str) -> str:
    """Save a plain text or markdown file that the user can open.

    Args:
        filename: File name, e.g. "notes.md" or "summary.txt".
        content: The file contents.

    Returns:
        A confirmation message with the saved path.
    """
    suffix = Path(filename).suffix or ".txt"
    path = _safe_output_path(filename, suffix)
    path.write_text(content, encoding="utf-8")
    return f"Saved to {path}. The user can download it from the Files panel."


UPLOAD_ROOT = Path.cwd() / "uploads"
MAX_EXTRACTED_CHARS = 20000

# Uploads are stored per conversation. A flat shared directory meant a new chat
# could see - and answer from - a document uploaded in a previous one.
_upload_scope: contextvars.ContextVar[str | None] = contextvars.ContextVar(
    "upload_scope", default=None
)


def set_upload_scope(thread_id: str | None) -> None:
    """Point the upload tools at one conversation's files."""
    _upload_scope.set(thread_id)


def upload_dir(thread_id: str | None = None) -> Path:
    """The upload directory for a conversation."""
    thread = thread_id or _upload_scope.get()
    if not thread:
        return UPLOAD_ROOT / "_unscoped"
    # The thread id comes from the browser, so keep it to a safe directory name.
    safe = re.sub(r"[^A-Za-z0-9._-]", "_", thread)[:80] or "_unscoped"
    return UPLOAD_ROOT / safe


def _truncate(text: str) -> str:
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS] + f"\n\n... (truncated at {MAX_EXTRACTED_CHARS} characters)"


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [f"--- page {i} ---\n{(page.extract_text() or '').strip()}" for i, page in enumerate(reader.pages, 1)]
    return "\n\n".join(pages)


def _read_docx(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for sheet in workbook.worksheets:
        out.append(f"--- sheet: {sheet.title} ---")
        for row in sheet.iter_rows(values_only=True):
            if any(cell is not None for cell in row):
                out.append(" | ".join("" if cell is None else str(cell) for cell in row))
    workbook.close()
    return "\n".join(out)


_READERS = {".pdf": _read_pdf, ".docx": _read_docx, ".xlsx": _read_xlsx, ".xlsm": _read_xlsx}

# Images are not read as text: they go to the model as image content blocks
# instead, since Qwen3.5 is multimodal. These other binaries cannot be read at
# all, so say so plainly rather than returning decoded noise.
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".svg", ".heic"}
BINARY_SUFFIXES = {".zip", ".exe", ".dll", ".mp4", ".mp3", ".mov", ".pptx", ".doc", ".xls"}


ZERO_WIDTH = "​‌‍﻿"


def _match_key(name: str) -> str:
    """A comparison key that survives a model retyping the filename.

    Small models re-tokenise non-ASCII filenames and reliably reinsert stray
    spaces around CJK runs ("Brief_2 頁" for "Brief_2頁"), so an exact match
    sends them into a retry loop they cannot escape. Compare with whitespace,
    zero-width characters, Unicode form and case all ignored.
    """
    name = unicodedata.normalize("NFKC", name)
    stripped = [c for c in name if not c.isspace() and c not in ZERO_WIDTH]
    return "".join(stripped).casefold()


def uploaded_files(thread_id: str | None = None) -> list[Path]:
    """This conversation's uploads, in the order their numbers refer to."""
    directory = upload_dir(thread_id)
    if not directory.exists():
        return []
    return sorted((p for p in directory.iterdir() if p.is_file()), key=lambda p: p.name)


def _resolve_upload(filename: str) -> Path | None:
    """Find the upload the model meant, tolerating an imperfectly typed name."""
    files = uploaded_files()
    if not files:
        return None

    wanted = (filename or "").strip()
    directory = upload_dir()

    # A 1-based index, which is what the agent is told to prefer.
    if wanted.isdigit():
        index = int(wanted)
        return files[index - 1] if 1 <= index <= len(files) else None

    exact = directory / Path(wanted).name
    if exact.is_file():
        return exact

    key = _match_key(Path(wanted).name)
    matches = [f for f in files if _match_key(f.name) == key]
    if len(matches) == 1:
        return matches[0]

    partial = [f for f in files if key and (key in _match_key(f.name) or _match_key(f.name) in key)]
    if len(partial) == 1:
        return partial[0]

    # Last resort: a single upload of the same type. Requiring the extension to
    # match matters — without it, a request for a file that was never uploaded
    # silently returns the one that was, and the model reports on the wrong
    # document rather than saying it could not find anything.
    suffix = Path(wanted).suffix.lower()
    if len(files) == 1 and suffix and files[0].suffix.lower() == suffix:
        return files[0]
    return None


def upload_listing(thread_id: str | None = None) -> str:
    """The numbered listing. The server builds its prompt from this same
    function, so the numbers the agent is shown always match the numbers
    `read_document` resolves."""
    files = uploaded_files(thread_id)
    if not files:
        return "none"
    return "; ".join(f"{i}. {f.name}" for i, f in enumerate(files, 1))


def _upload_listing() -> str:
    return upload_listing()


def read_document(filename: str) -> str:
    """Read a file the user uploaded, and return its text.

    Supports .pdf, .docx, .xlsx, .csv, .txt and .md.

    Args:
        filename: The number of the attachment as shown in the message
            (e.g. "1"), which is the most reliable way to refer to it, or the
            file name.

    Returns:
        The file's text content, truncated if very long.
    """
    path = _resolve_upload(filename)
    if path is None:
        return (
            f"No uploaded file matches {filename!r}. "
            f"Available uploads: {_upload_listing()}. "
            f"Call read_document with the NUMBER instead, e.g. read_document('1'). "
            f"Do not retry the same filename."
        )

    suffix = path.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        return (
            f"{path.name!r} is an image and has already been attached to this "
            f"conversation visually - look at it directly and describe what you "
            f"see. read_document only extracts text, so there is nothing to read "
            f"here. Do not read a different file instead."
        )
    if suffix in BINARY_SUFFIXES:
        return f"{path.name!r} is a {suffix} file, which cannot be read as text."

    reader = _READERS.get(suffix)
    try:
        text = reader(path) if reader else path.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:
        return f"Could not read {filename!r}: {type(exc).__name__}: {exc}"

    return _truncate(text) if text.strip() else f"{filename!r} contains no extractable text."


def list_uploads() -> str:
    """List the files the user has uploaded, numbered for use with read_document."""
    listing = _upload_listing()
    if listing == "none":
        return "No files uploaded."
    return f"{listing}\n\nRead one with its number, e.g. read_document('1')."


# Images are sent exactly as uploaded - no resizing, no re-encoding - so the
# model sees the full detail. Note that an attached image stays in the
# conversation history and is re-sent on every model call in that thread, so a
# large photo costs its tokens repeatedly.
IMAGE_MIME = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".tiff": "image/tiff",
    ".heic": "image/heic",
    ".svg": "image/svg+xml",
}


def encode_image(path: Path) -> tuple[str, str]:
    """Return (mime_type, base64 data) for an image, byte-for-byte as uploaded."""
    import base64

    mime = IMAGE_MIME.get(path.suffix.lower(), "application/octet-stream")
    return mime, base64.b64encode(path.read_bytes()).decode("ascii")


def image_attachments(names: list[str], thread_id: str | None = None) -> list[dict]:
    """Build the model content blocks for any images among these attachments.

    langchain converts this standard block into the OpenAI `image_url` data-URL
    form that vLLM expects.
    """
    blocks: list[dict] = []
    directory = upload_dir(thread_id)
    for name in names:
        path = directory / Path(name).name
        if not path.is_file() or path.suffix.lower() not in IMAGE_SUFFIXES:
            continue
        try:
            mime, data = encode_image(path)
        except Exception:
            continue  # unreadable image: fall back to treating it as a plain file
        blocks.append(
            {"type": "image", "source_type": "base64", "mime_type": mime, "data": data}
        )
    return blocks


def is_image(name: str) -> bool:
    return Path(name).suffix.lower() in IMAGE_SUFFIXES
